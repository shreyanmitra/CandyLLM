"""
🍭 CandyLLM Multimodal Support
Vision, audio, and document processing capabilities
"""

import os
import base64
import mimetypes
from abc import ABC, abstractmethod
from typing import Dict, List, Any, Optional, Union, BinaryIO
from dataclasses import dataclass
from pathlib import Path
import tempfile
import asyncio

@dataclass
class MediaContent:
    """Represents multimodal content"""
    content_type: str  # "image", "audio", "video", "document"
    mime_type: str
    data: Union[bytes, str]  # bytes for binary, str for text/base64
    metadata: Dict[str, Any]
    
    @property
    def is_binary(self) -> bool:
        return isinstance(self.data, bytes)
    
    def to_base64(self) -> str:
        """Convert to base64 string"""
        if self.is_binary:
            return base64.b64encode(self.data).decode('utf-8')
        return self.data
    
    def save_to_file(self, path: str) -> str:
        """Save content to file"""
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        
        if self.is_binary:
            with open(path, 'wb') as f:
                f.write(self.data)
        else:
            with open(path, 'w') as f:
                f.write(self.data)
        
        return path

class MediaProcessor(ABC):
    """Abstract base class for media processors"""
    
    @abstractmethod
    def can_process(self, content: MediaContent) -> bool:
        """Check if this processor can handle the content"""
        pass
    
    @abstractmethod
    async def process(self, content: MediaContent) -> Dict[str, Any]:
        """Process the media content"""
        pass

class ImageProcessor(MediaProcessor):
    """Image processing and analysis"""
    
    def __init__(self):
        self.supported_formats = {
            'image/jpeg', 'image/png', 'image/gif', 'image/bmp', 
            'image/webp', 'image/tiff', 'image/svg+xml'
        }
    
    def can_process(self, content: MediaContent) -> bool:
        return (content.content_type == "image" and 
                content.mime_type in self.supported_formats)
    
    async def process(self, content: MediaContent) -> Dict[str, Any]:
        """Process image content"""
        try:
            # Try to use PIL for image analysis
            from PIL import Image
            import io
            
            if content.is_binary:
                image = Image.open(io.BytesIO(content.data))
            else:
                # Assume base64 encoded
                image_data = base64.b64decode(content.data)
                image = Image.open(io.BytesIO(image_data))
            
            # Extract basic information
            info = {
                "format": image.format,
                "mode": image.mode,
                "size": image.size,
                "width": image.width,
                "height": image.height,
                "has_transparency": image.mode in ('RGBA', 'LA', 'P'),
                "file_size": len(content.data) if content.is_binary else len(base64.b64decode(content.data))
            }
            
            # Try to extract EXIF data
            try:
                from PIL.ExifTags import TAGS
                exif_data = {}
                if hasattr(image, '_getexif') and image._getexif():
                    exif = image._getexif()
                    for tag_id, value in exif.items():
                        tag = TAGS.get(tag_id, tag_id)
                        exif_data[tag] = value
                info["exif"] = exif_data
            except:
                info["exif"] = {}
            
            # Analyze colors (simplified)
            try:
                colors = image.getcolors(maxcolors=256)
                if colors:
                    dominant_color = max(colors, key=lambda x: x[0])[1]
                    info["dominant_color"] = dominant_color
            except:
                pass
            
            return {
                "status": "success",
                "type": "image_analysis",
                "info": info,
                "description": f"{info['format']} image ({info['width']}x{info['height']})",
                "processable_by_llm": True
            }
            
        except ImportError:
            return {
                "status": "error",
                "error": "PIL (Pillow) not installed. Run: pip install Pillow",
                "processable_by_llm": False
            }
        except Exception as e:
            return {
                "status": "error",
                "error": str(e),
                "processable_by_llm": False
            }

class AudioProcessor(MediaProcessor):
    """Audio processing and transcription"""
    
    def __init__(self):
        self.supported_formats = {
            'audio/wav', 'audio/mp3', 'audio/flac', 'audio/ogg',
            'audio/m4a', 'audio/webm', 'audio/mpeg'
        }
    
    def can_process(self, content: MediaContent) -> bool:
        return (content.content_type == "audio" and 
                content.mime_type in self.supported_formats)
    
    async def process(self, content: MediaContent) -> Dict[str, Any]:
        """Process audio content"""
        try:
            # Try to use librosa for audio analysis
            import librosa
            import numpy as np
            import io
            
            # Save to temporary file for librosa
            with tempfile.NamedTemporaryFile(suffix='.wav', delete=False) as temp_file:
                if content.is_binary:
                    temp_file.write(content.data)
                else:
                    audio_data = base64.b64decode(content.data)
                    temp_file.write(audio_data)
                temp_path = temp_file.name
            
            try:
                # Load audio
                y, sr = librosa.load(temp_path)
                
                # Extract features
                duration = len(y) / sr
                tempo, beats = librosa.beat.beat_track(y=y, sr=sr)
                spectral_centroids = librosa.feature.spectral_centroid(y=y, sr=sr)[0]
                zero_crossings = librosa.feature.zero_crossing_rate(y)[0]
                
                info = {
                    "duration": float(duration),
                    "sample_rate": int(sr),
                    "channels": 1 if y.ndim == 1 else y.shape[0],
                    "tempo": float(tempo),
                    "avg_spectral_centroid": float(np.mean(spectral_centroids)),
                    "avg_zero_crossing_rate": float(np.mean(zero_crossings)),
                    "file_size": len(content.data) if content.is_binary else len(base64.b64decode(content.data))
                }
                
                # Try transcription with whisper if available
                transcription = await self._try_transcription(temp_path)
                if transcription:
                    info["transcription"] = transcription
                
                return {
                    "status": "success",
                    "type": "audio_analysis",
                    "info": info,
                    "description": f"Audio file ({duration:.1f}s, {sr}Hz)",
                    "processable_by_llm": bool(transcription)
                }
                
            finally:
                # Clean up temp file
                os.unlink(temp_path)
                
        except ImportError:
            return {
                "status": "error",
                "error": "librosa not installed. Run: pip install librosa",
                "processable_by_llm": False
            }
        except Exception as e:
            return {
                "status": "error",
                "error": str(e),
                "processable_by_llm": False
            }
    
    async def _try_transcription(self, audio_path: str) -> Optional[str]:
        """Try to transcribe audio using Whisper"""
        try:
            import whisper
            
            model = whisper.load_model("base")
            result = model.transcribe(audio_path)
            return result["text"]
            
        except ImportError:
            return None
        except Exception:
            return None

class DocumentProcessor(MediaProcessor):
    """Document processing and text extraction"""
    
    def __init__(self):
        self.supported_formats = {
            'application/pdf',
            'text/plain', 'text/markdown', 'text/html',
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',  # docx
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',  # xlsx
            'application/json', 'application/xml'
        }
    
    def can_process(self, content: MediaContent) -> bool:
        return (content.content_type == "document" and 
                content.mime_type in self.supported_formats)
    
    async def process(self, content: MediaContent) -> Dict[str, Any]:
        """Process document content"""
        try:
            if content.mime_type == 'application/pdf':
                return await self._process_pdf(content)
            elif content.mime_type.startswith('text/'):
                return await self._process_text(content)
            elif 'word' in content.mime_type:
                return await self._process_docx(content)
            elif 'sheet' in content.mime_type:
                return await self._process_xlsx(content)
            else:
                return {
                    "status": "error",
                    "error": f"Unsupported document type: {content.mime_type}",
                    "processable_by_llm": False
                }
                
        except Exception as e:
            return {
                "status": "error",
                "error": str(e),
                "processable_by_llm": False
            }
    
    async def _process_pdf(self, content: MediaContent) -> Dict[str, Any]:
        """Process PDF document"""
        try:
            import PyPDF2
            import io
            
            if content.is_binary:
                pdf_data = content.data
            else:
                pdf_data = base64.b64decode(content.data)
            
            reader = PyPDF2.PdfReader(io.BytesIO(pdf_data))
            
            text_content = ""
            for page in reader.pages:
                text_content += page.extract_text() + "\n"
            
            info = {
                "pages": len(reader.pages),
                "text_length": len(text_content),
                "file_size": len(pdf_data),
                "text_content": text_content
            }
            
            # Try to extract metadata
            if reader.metadata:
                info["metadata"] = {
                    "title": reader.metadata.get('/Title', ''),
                    "author": reader.metadata.get('/Author', ''),
                    "creator": reader.metadata.get('/Creator', ''),
                    "producer": reader.metadata.get('/Producer', ''),
                    "subject": reader.metadata.get('/Subject', '')
                }
            
            return {
                "status": "success",
                "type": "pdf_analysis",
                "info": info,
                "description": f"PDF document ({info['pages']} pages, {info['text_length']} characters)",
                "processable_by_llm": True
            }
            
        except ImportError:
            return {
                "status": "error",
                "error": "PyPDF2 not installed. Run: pip install PyPDF2",
                "processable_by_llm": False
            }
    
    async def _process_text(self, content: MediaContent) -> Dict[str, Any]:
        """Process text document"""
        if content.is_binary:
            text_content = content.data.decode('utf-8', errors='ignore')
        else:
            text_content = content.data
        
        info = {
            "text_length": len(text_content),
            "line_count": text_content.count('\n') + 1,
            "word_count": len(text_content.split()),
            "text_content": text_content
        }
        
        return {
            "status": "success",
            "type": "text_analysis",
            "info": info,
            "description": f"Text document ({info['word_count']} words, {info['line_count']} lines)",
            "processable_by_llm": True
        }
    
    async def _process_docx(self, content: MediaContent) -> Dict[str, Any]:
        """Process Word document"""
        try:
            from docx import Document
            import io
            
            if content.is_binary:
                doc_data = content.data
            else:
                doc_data = base64.b64decode(content.data)
            
            doc = Document(io.BytesIO(doc_data))
            
            text_content = ""
            for paragraph in doc.paragraphs:
                text_content += paragraph.text + "\n"
            
            info = {
                "paragraphs": len(doc.paragraphs),
                "text_length": len(text_content),
                "file_size": len(doc_data),
                "text_content": text_content
            }
            
            return {
                "status": "success",
                "type": "docx_analysis",
                "info": info,
                "description": f"Word document ({info['paragraphs']} paragraphs, {info['text_length']} characters)",
                "processable_by_llm": True
            }
            
        except ImportError:
            return {
                "status": "error",
                "error": "python-docx not installed. Run: pip install python-docx",
                "processable_by_llm": False
            }
    
    async def _process_xlsx(self, content: MediaContent) -> Dict[str, Any]:
        """Process Excel spreadsheet"""
        try:
            import pandas as pd
            import io
            
            if content.is_binary:
                excel_data = content.data
            else:
                excel_data = base64.b64decode(content.data)
            
            # Read all sheets
            excel_file = pd.ExcelFile(io.BytesIO(excel_data))
            sheets_data = {}
            
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                sheets_data[sheet_name] = {
                    "rows": len(df),
                    "columns": len(df.columns),
                    "column_names": df.columns.tolist(),
                    "data_preview": df.head().to_dict('records')
                }
            
            # Create text summary
            text_content = f"Excel file with {len(sheets_data)} sheets:\n"
            for sheet_name, sheet_info in sheets_data.items():
                text_content += f"\nSheet '{sheet_name}': {sheet_info['rows']} rows, {sheet_info['columns']} columns\n"
                text_content += f"Columns: {', '.join(sheet_info['column_names'])}\n"
            
            info = {
                "sheets": sheets_data,
                "total_sheets": len(sheets_data),
                "file_size": len(excel_data),
                "text_content": text_content
            }
            
            return {
                "status": "success",
                "type": "xlsx_analysis",
                "info": info,
                "description": f"Excel file ({info['total_sheets']} sheets)",
                "processable_by_llm": True
            }
            
        except ImportError:
            return {
                "status": "error",
                "error": "pandas not installed. Run: pip install pandas openpyxl",
                "processable_by_llm": False
            }

class MultimodalManager:
    """
    Central manager for multimodal content processing
    """
    
    def __init__(self):
        self.processors = [
            ImageProcessor(),
            AudioProcessor(),
            DocumentProcessor()
        ]
    
    def load_from_file(self, file_path: str) -> MediaContent:
        """Load content from file"""
        path = Path(file_path)
        
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        # Determine content type and MIME type
        mime_type, _ = mimetypes.guess_type(file_path)
        if not mime_type:
            mime_type = "application/octet-stream"
        
        content_type = self._determine_content_type(mime_type)
        
        # Read file content
        if content_type in ["image", "audio"] or mime_type.startswith("application/"):
            # Binary content
            with open(file_path, 'rb') as f:
                data = f.read()
        else:
            # Text content
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                data = f.read()
        
        metadata = {
            "file_path": str(path.absolute()),
            "file_name": path.name,
            "file_size": path.stat().st_size,
            "created_time": path.stat().st_ctime,
            "modified_time": path.stat().st_mtime
        }
        
        return MediaContent(
            content_type=content_type,
            mime_type=mime_type,
            data=data,
            metadata=metadata
        )
    
    def load_from_base64(self, base64_data: str, mime_type: str) -> MediaContent:
        """Load content from base64 string"""
        content_type = self._determine_content_type(mime_type)
        
        return MediaContent(
            content_type=content_type,
            mime_type=mime_type,
            data=base64_data,
            metadata={"source": "base64"}
        )
    
    def load_from_bytes(self, data: bytes, mime_type: str) -> MediaContent:
        """Load content from bytes"""
        content_type = self._determine_content_type(mime_type)
        
        return MediaContent(
            content_type=content_type,
            mime_type=mime_type,
            data=data,
            metadata={"source": "bytes", "size": len(data)}
        )
    
    def _determine_content_type(self, mime_type: str) -> str:
        """Determine content type from MIME type"""
        if mime_type.startswith("image/"):
            return "image"
        elif mime_type.startswith("audio/"):
            return "audio"
        elif mime_type.startswith("video/"):
            return "video"
        elif mime_type.startswith("text/") or mime_type in [
            "application/pdf", "application/json", "application/xml",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ]:
            return "document"
        else:
            return "unknown"
    
    async def process_content(self, content: MediaContent) -> Dict[str, Any]:
        """Process multimodal content"""
        # Find appropriate processor
        processor = None
        for proc in self.processors:
            if proc.can_process(content):
                processor = proc
                break
        
        if not processor:
            return {
                "status": "error",
                "error": f"No processor available for {content.content_type} ({content.mime_type})",
                "processable_by_llm": False
            }
        
        # Process content
        result = await processor.process(content)
        
        # Add general metadata
        result.update({
            "content_type": content.content_type,
            "mime_type": content.mime_type,
            "metadata": content.metadata
        })
        
        return result
    
    async def process_file(self, file_path: str) -> Dict[str, Any]:
        """Process a file"""
        try:
            content = self.load_from_file(file_path)
            return await self.process_content(content)
        except Exception as e:
            return {
                "status": "error",
                "error": str(e),
                "processable_by_llm": False
            }
    
    def get_supported_formats(self) -> Dict[str, List[str]]:
        """Get list of supported formats by content type"""
        formats = {}
        
        for processor in self.processors:
            if hasattr(processor, 'supported_formats'):
                content_types = set()
                for mime_type in processor.supported_formats:
                    content_type = self._determine_content_type(mime_type)
                    content_types.add(content_type)
                
                for content_type in content_types:
                    if content_type not in formats:
                        formats[content_type] = []
                    formats[content_type].extend([
                        mime for mime in processor.supported_formats
                        if self._determine_content_type(mime) == content_type
                    ])
        
        return formats

# Global multimodal manager instance
_multimodal_manager = None

def get_multimodal_manager() -> MultimodalManager:
    """Get global multimodal manager instance"""
    global _multimodal_manager
    if _multimodal_manager is None:
        _multimodal_manager = MultimodalManager()
    return _multimodal_manager

async def process_file(file_path: str) -> Dict[str, Any]:
    """Convenience function to process a file"""
    manager = get_multimodal_manager()
    return await manager.process_file(file_path)

async def process_base64(base64_data: str, mime_type: str) -> Dict[str, Any]:
    """Convenience function to process base64 data"""
    manager = get_multimodal_manager()
    content = manager.load_from_base64(base64_data, mime_type)
    return await manager.process_content(content)
